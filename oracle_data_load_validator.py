#!/usr/bin/env python3
"""Generic mapping-driven Excel data-load validator.

The workbook must contain a sheet named "Mapping-Sheet" with these headers:
    Source Tab | Source Column Name | Key Sequence | Dest Tab | Dest Column Name

Key Sequence is populated only for key columns (1, 2, 3, ...). Blank means the
column is compared but is not part of the record key.

Examples:
    python mapping_sheet_validator.py "input.xlsx"
    python mapping_sheet_validator.py "input.xlsx" --output "report.xlsx"
    python mapping_sheet_validator.py "C:\\Data\\ToValidate" --output-dir "C:\\Data\\Reports"

Dependencies:
    py -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAPPING_SHEET = "Mapping-Sheet"
OUTPUT_SUFFIX = "_Validation_Report.xlsx"

HEADER_ALIASES = {
    "source_tab": {"source tab", "source sheet", "source worksheet"},
    "source_column": {"source column name", "source column", "source field"},
    "key_sequence": {"key sequence", "key seq", "key order", "key position"},
    "dest_tab": {"dest tab", "destination tab", "destination sheet", "destination worksheet"},
    "dest_column": {"dest column name", "destination column name", "dest column", "destination column"},
}


@dataclass(frozen=True)
class MappingRow:
    source_sheet: str
    source_column: str
    key_sequence: int | None
    destination_sheet: str
    destination_column: str
    mapping_row: int


def normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def normalized_value(value: Any, ignore_case: bool = False) -> str:
    """Normalize equivalent Excel representations without changing identifiers."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()
    if not text:
        return ""

    # Normalize common date strings so Excel dates and report text dates compare equally.
    for pattern in ("%Y/%m/%d", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            parsed = datetime.strptime(text, pattern)
            text = parsed.date().isoformat() if parsed.time().isoformat() == "00:00:00" else parsed.isoformat()
            break
        except ValueError:
            pass

    # Treat 1, 1.0 and 1.00 as equal, while retaining integer IDs such as 00123.
    numeric_pattern = r"[-+]?\d+(?:\.\d+)?"
    has_integer_leading_zero = (
        len(text) > 1 and text[0] == "0" and text[1].isdigit() and "." not in text
    )
    if re.fullmatch(numeric_pattern, text) and not has_integer_leading_zero:
        try:
            number = Decimal(text)
            if number == number.to_integral():
                text = str(number.quantize(Decimal("1")))
            else:
                text = format(number.normalize(), "f")
        except InvalidOperation:
            pass

    return text.casefold() if ignore_case else text


def find_mapping_columns(ws) -> tuple[int, dict[str, int]]:
    """Find the Mapping-Sheet header row and required column indexes."""
    for row_number in range(1, min(ws.max_row, 50) + 1):
        seen = {
            normalized_header(ws.cell(row_number, column).value): column
            for column in range(1, ws.max_column + 1)
            if ws.cell(row_number, column).value is not None
        }
        resolved: dict[str, int] = {}
        for logical_name, aliases in HEADER_ALIASES.items():
            match = next((seen[alias] for alias in aliases if alias in seen), None)
            if match is not None:
                resolved[logical_name] = match
        if len(resolved) == len(HEADER_ALIASES):
            return row_number, resolved

    expected = ", ".join(sorted(next(iter(v)) for v in HEADER_ALIASES.values()))
    raise ValueError(
        f"Could not find all required headers in '{MAPPING_SHEET}'. Expected: {expected}"
    )


def parse_key_sequence(value: Any, excel_row: int) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"Mapping-Sheet row {excel_row}: invalid Key Sequence {value!r}") from exc
    if number != number.to_integral() or number <= 0:
        raise ValueError(
            f"Mapping-Sheet row {excel_row}: Key Sequence must be a positive integer"
        )
    return int(number)


def read_mappings(wb) -> dict[tuple[str, str], list[MappingRow]]:
    if MAPPING_SHEET not in wb.sheetnames:
        raise ValueError(f"Required worksheet '{MAPPING_SHEET}' was not found")

    ws = wb[MAPPING_SHEET]
    header_row, columns = find_mapping_columns(ws)
    grouped: dict[tuple[str, str], list[MappingRow]] = defaultdict(list)

    for row_number in range(header_row + 1, ws.max_row + 1):
        source_sheet = ws.cell(row_number, columns["source_tab"]).value
        source_column = ws.cell(row_number, columns["source_column"]).value
        destination_sheet = ws.cell(row_number, columns["dest_tab"]).value
        destination_column = ws.cell(row_number, columns["dest_column"]).value
        key_raw = ws.cell(row_number, columns["key_sequence"]).value

        values = [source_sheet, source_column, destination_sheet, destination_column]
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        if any(v is None or str(v).strip() == "" for v in values):
            raise ValueError(
                f"Mapping-Sheet row {row_number}: source/destination sheet and column values are required"
            )

        item = MappingRow(
            source_sheet=str(source_sheet),
            source_column=str(source_column),
            key_sequence=parse_key_sequence(key_raw, row_number),
            destination_sheet=str(destination_sheet),
            destination_column=str(destination_column),
            mapping_row=row_number,
        )
        grouped[(item.source_sheet, item.destination_sheet)].append(item)

    if not grouped:
        raise ValueError("Mapping-Sheet has no usable mapping rows")

    for (source_sheet, destination_sheet), rows in grouped.items():
        if source_sheet not in wb.sheetnames:
            raise ValueError(f"Mapped source worksheet not found: {source_sheet!r}")
        if destination_sheet not in wb.sheetnames:
            raise ValueError(f"Mapped destination worksheet not found: {destination_sheet!r}")

        key_rows = [row for row in rows if row.key_sequence is not None]
        if not key_rows:
            raise ValueError(
                f"No Key Sequence is defined for mapping {source_sheet!r} -> {destination_sheet!r}"
            )
        sequences = [row.key_sequence for row in key_rows]
        expected = list(range(1, len(sequences) + 1))
        if sorted(sequences) != expected:
            raise ValueError(
                f"Key Sequence for {source_sheet!r} -> {destination_sheet!r} must be unique and contiguous: {expected}"
            )

        source_names = [normalized_header(row.source_column) for row in rows]
        destination_names = [normalized_header(row.destination_column) for row in rows]
        if len(source_names) != len(set(source_names)):
            raise ValueError(f"Duplicate source-column mapping in {source_sheet!r}")
        if len(destination_names) != len(set(destination_names)):
            raise ValueError(f"Duplicate destination-column mapping in {destination_sheet!r}")

    return grouped


def locate_data_header(ws, expected_columns: Iterable[str]) -> tuple[int, dict[str, int]]:
    expected = list(expected_columns)
    expected_normalized = {normalized_header(column) for column in expected}
    best_score = -1
    best_row = None
    best_map: dict[str, int] = {}

    for row_number in range(1, min(ws.max_row, 100) + 1):
        column_map = {
            normalized_header(ws.cell(row_number, column).value): column
            for column in range(1, ws.max_column + 1)
            if ws.cell(row_number, column).value is not None
        }
        score = len(expected_normalized.intersection(column_map))
        if score > best_score:
            best_score, best_row, best_map = score, row_number, column_map

    if best_score != len(expected_normalized):
        missing = [column for column in expected if normalized_header(column) not in best_map]
        raise ValueError(f"Worksheet {ws.title!r}: mapped headers not found: {missing}")
    return int(best_row), best_map


def read_records(ws, header_row: int, columns: dict[str, int], ignore_case: bool):
    records = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        record = {name: ws.cell(row_number, column).value for name, column in columns.items()}
        if any(normalized_value(value, ignore_case) != "" for value in record.values()):
            records.append((row_number, record))
    return records


def has_complete_key(record: dict[str, Any], key_fields: list[str], ignore_case: bool) -> bool:
    return all(normalized_value(record[field], ignore_case) != "" for field in key_fields)


def choose_fallback_key(source_records, destination_records, source_fields, ignore_case):
    """Choose a safe alternate field only when configured keys have no overlap.

    A fallback field must be nonblank, unique on both sides, and have the same
    normalized value set. This prevents generated source IDs from turning every
    row into a false Missing/Unexpected exception while avoiding fuzzy matching.
    """
    candidates = []
    for field in source_fields:
        source_values = [normalized_value(record[field], ignore_case) for _, record in source_records]
        destination_values = [normalized_value(record[field], ignore_case) for _, record in destination_records]
        if not source_values or not destination_values:
            continue
        if "" in source_values or "" in destination_values:
            continue
        if len(source_values) != len(set(source_values)) or len(destination_values) != len(set(destination_values)):
            continue
        if set(source_values) == set(destination_values):
            candidates.append(field)
    return [candidates[0]] if candidates else []


def make_key(record: dict[str, Any], key_fields: list[str], ignore_case: bool) -> tuple[str, ...]:
    return tuple(normalized_value(record[field], ignore_case) for field in key_fields)


def signature(record: dict[str, Any], fields: list[str], ignore_case: bool) -> tuple[str, ...]:
    return tuple(normalized_value(record[field], ignore_case) for field in fields)


def display_key(key: tuple[str, ...]) -> str:
    return " | ".join(key)


def style_report_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    for column in range(1, ws.max_column + 1):
        sample = [str(ws.cell(row, column).value or "") for row in range(1, min(ws.max_row, 250) + 1)]
        width = min(48, max(12, max((len(value) for value in sample), default=10) + 2))
        ws.column_dimensions[get_column_letter(column)].width = width



def normalize_xlsx_font_order(path: Path) -> None:
    """Normalize styles.xml using only Python standard-library modules."""
    temp_path = path.with_suffix(path.suffix + ".tmp")
    order = {
        "b": 0, "i": 1, "strike": 2, "outline": 3, "shadow": 4,
        "condense": 5, "extend": 6, "sz": 7, "color": 8,
        "name": 9, "family": 10, "charset": 11, "scheme": 12,
        "u": 13, "vertAlign": 14,
    }
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)
    with zipfile.ZipFile(path, "r") as source_zip, zipfile.ZipFile(
        temp_path, "w", zipfile.ZIP_DEFLATED
    ) as target_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename == "xl/styles.xml":
                root = ET.fromstring(data)
                fonts = root.find(f"{{{namespace}}}fonts")
                if fonts is not None:
                    for font in fonts:
                        children = list(font)
                        for child in children:
                            font.remove(child)
                        children.sort(
                            key=lambda element: order.get(element.tag.rsplit("}", 1)[-1], 99)
                        )
                        for child in children:
                            font.append(child)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target_zip.writestr(item, data)
    os.replace(temp_path, path)

def append_table(workbook, name: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = workbook.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_report_sheet(ws)


def validate_workbook(input_path: Path, output_path: Path, ignore_case: bool = False) -> None:
    wb = load_workbook(input_path, data_only=False, read_only=False)
    grouped_mappings = read_mappings(wb)

    reconciliation = []
    missing_records = []
    data_mismatches = []
    duplicate_records = []
    unexpected_records = []
    metadata_warnings = []

    for (source_sheet, destination_sheet), mappings in grouped_mappings.items():
        source_ws, destination_ws = wb[source_sheet], wb[destination_sheet]
        source_fields = [row.source_column for row in mappings]
        key_fields = [
            row.source_column
            for row in sorted(
                (row for row in mappings if row.key_sequence is not None),
                key=lambda row: row.key_sequence,
            )
        ]

        source_header_row, source_header_map = locate_data_header(source_ws, source_fields)
        destination_header_row, destination_header_map = locate_data_header(
            destination_ws, [row.destination_column for row in mappings]
        )
        source_columns = {
            row.source_column: source_header_map[normalized_header(row.source_column)]
            for row in mappings
        }
        # Alias destination records using source field names so comparisons are mapping-driven.
        destination_columns = {
            row.source_column: destination_header_map[normalized_header(row.destination_column)]
            for row in mappings
        }

        source_records = read_records(source_ws, source_header_row, source_columns, ignore_case)
        destination_records = read_records(
            destination_ws, destination_header_row, destination_columns, ignore_case
        )

        # Exclude report totals, footers and notes. A data row must have every
        # configured key component populated.
        source_records = [item for item in source_records if has_complete_key(item[1], key_fields, ignore_case)]
        destination_records = [item for item in destination_records if has_complete_key(item[1], key_fields, ignore_case)]

        effective_key_fields = list(key_fields)
        source_groups: dict[tuple[str, ...], list[tuple[int, dict[str, Any]]]] = defaultdict(list)
        destination_groups: dict[tuple[str, ...], list[tuple[int, dict[str, Any]]]] = defaultdict(list)
        for item in source_records:
            source_groups[make_key(item[1], effective_key_fields, ignore_case)].append(item)
        for item in destination_records:
            destination_groups[make_key(item[1], effective_key_fields, ignore_case)].append(item)

        # If configured keys have zero overlap, they likely contain IDs generated
        # by the destination system. Use only a deterministic, unique, exact
        # fallback field. Never use fuzzy matching.
        if source_records and destination_records and not (set(source_groups) & set(destination_groups)):
            fallback = choose_fallback_key(source_records, destination_records, source_fields, ignore_case)
            if fallback:
                metadata_warnings.append([
                    source_sheet, destination_sheet, ", ".join(key_fields), ", ".join(fallback),
                    "Configured key values have no source/destination overlap; a unique exact fallback field was used for pairing. Review Key Sequence and generated-ID mappings."
                ])
                effective_key_fields = fallback
                source_groups = defaultdict(list)
                destination_groups = defaultdict(list)
                for item in source_records:
                    source_groups[make_key(item[1], effective_key_fields, ignore_case)].append(item)
                for item in destination_records:
                    destination_groups[make_key(item[1], effective_key_fields, ignore_case)].append(item)

        for side, groups in (("Source", source_groups), ("Destination", destination_groups)):
            for key, rows in groups.items():
                if len(rows) > 1:
                    for row_number, _ in rows:
                        duplicate_records.append([
                            source_sheet, destination_sheet, display_key(key), row_number,
                            side, len(rows), "Duplicate Record",
                        ])

        matched = missing_count = unexpected_count = mismatch_record_count = 0

        for key in set(source_groups) | set(destination_groups):
            source_group = list(source_groups.get(key, []))
            destination_group = list(destination_groups.get(key, []))

            # Exact full-record matching first, including repeated-key groups.
            destination_by_signature: dict[tuple[str, ...], list[tuple[int, dict[str, Any]]]] = defaultdict(list)
            for item in destination_group:
                destination_by_signature[signature(item[1], source_fields, ignore_case)].append(item)

            source_remaining = []
            for item in source_group:
                row_signature = signature(item[1], source_fields, ignore_case)
                if destination_by_signature[row_signature]:
                    destination_by_signature[row_signature].pop()
                    matched += 1
                else:
                    source_remaining.append(item)
            destination_remaining = [item for bucket in destination_by_signature.values() for item in bucket]

            # Pair residual rows with the same key by minimum mapped-field differences.
            while source_remaining and destination_remaining:
                difference_count, source_index, destination_index = min(
                    (
                        sum(
                            normalized_value(source_record[field], ignore_case)
                            != normalized_value(destination_record[field], ignore_case)
                            for field in source_fields
                        ),
                        source_index,
                        destination_index,
                    )
                    for source_index, (_, source_record) in enumerate(source_remaining)
                    for destination_index, (_, destination_record) in enumerate(destination_remaining)
                )
                source_row, source_record = source_remaining.pop(source_index)
                destination_row, destination_record = destination_remaining.pop(destination_index)
                mismatch_record_count += 1

                for mapping in mappings:
                    source_value = source_record[mapping.source_column]
                    destination_value = destination_record[mapping.source_column]
                    if normalized_value(source_value, ignore_case) != normalized_value(
                        destination_value, ignore_case
                    ):
                        data_mismatches.append([
                            source_sheet, destination_sheet, display_key(key),
                            mapping.source_column, mapping.destination_column,
                            source_value, destination_value, "Data Mismatch",
                            source_row, destination_row,
                        ])

            for source_row, source_record in source_remaining:
                missing_count += 1
                missing_records.append([
                    source_sheet, destination_sheet, display_key(key), source_row,
                    "; ".join(f"{field}={source_record[field]}" for field in source_fields),
                    "Missing Record",
                ])
            for destination_row, destination_record in destination_remaining:
                unexpected_count += 1
                unexpected_records.append([
                    source_sheet, destination_sheet, display_key(key), destination_row,
                    "; ".join(f"{field}={destination_record[field]}" for field in source_fields),
                    "Unexpected Record",
                ])

        source_duplicate_rows = sum(
            len(rows) for rows in source_groups.values() if len(rows) > 1
        )
        destination_duplicate_rows = sum(
            len(rows) for rows in destination_groups.values() if len(rows) > 1
        )
        reconciliation.append([
            source_sheet, destination_sheet, ", ".join(effective_key_fields),
            len(source_records), len(destination_records), matched,
            missing_count, unexpected_count, mismatch_record_count,
            source_duplicate_rows, destination_duplicate_rows,
        ])

    total_source = sum(row[3] for row in reconciliation)
    total_destination = sum(row[4] for row in reconciliation)
    total_matched = sum(row[5] for row in reconciliation)
    total_missing = sum(row[6] for row in reconciliation)
    total_unexpected = sum(row[7] for row in reconciliation)
    total_mismatched = sum(row[8] for row in reconciliation)
    total_duplicate_rows = sum(row[9] + row[10] for row in reconciliation)
    status = "PASS" if not (
        total_missing or total_unexpected or total_mismatched or total_duplicate_rows
    ) else "FAIL"

    report = Workbook()
    report.remove(report.active)

    executive_rows = [
        ["Overall Status", status],
        ["Mappings Validated", len(reconciliation)],
        ["Total Source Records", total_source],
        ["Total Destination Records", total_destination],
        ["Matched Records", total_matched],
        ["Missing Records", total_missing],
        ["Unexpected Records", total_unexpected],
        ["Records with Data Mismatches", total_mismatched],
        ["Duplicate Rows Flagged", total_duplicate_rows],
    ]
    append_table(report, "Executive Summary", ["Metric", "Value"], executive_rows)

    reconciliation_headers = [
        "Source Worksheet", "Destination Worksheet", "Key Fields",
        "Total Source Records", "Total Destination Records", "Matched Records",
        "Missing Records", "Unexpected Records", "Records with Data Mismatches",
        "Source Duplicate Rows", "Destination Duplicate Rows",
    ]
    append_table(report, "Record Count Reconciliation", reconciliation_headers, reconciliation)
    append_table(
        report, "Missing Records",
        ["Source Worksheet", "Destination Worksheet", "Key Field Value(s)",
         "Source Row", "Source Record", "Issue Type"],
        missing_records,
    )
    append_table(
        report, "Data Mismatches",
        ["Source Worksheet", "Destination Worksheet", "Key Field Value(s)",
         "Source Column", "Destination Column", "Source Value", "Destination Value",
         "Issue Type", "Source Row", "Destination Row"],
        data_mismatches,
    )
    append_table(
        report, "Duplicate Records",
        ["Source Worksheet", "Destination Worksheet", "Key Field Value(s)",
         "Worksheet Row", "Side", "Duplicate Group Size", "Issue Type"],
        duplicate_records,
    )
    append_table(
        report, "Unexpected Records",
        ["Source Worksheet", "Destination Worksheet", "Key Field Value(s)",
         "Destination Row", "Destination Record", "Issue Type"],
        unexpected_records,
    )

    append_table(
        report, "Metadata Warnings",
        ["Source Worksheet", "Destination Worksheet", "Configured Key Fields",
         "Effective Pairing Fields", "Warning"],
        metadata_warnings,
    )

    mismatch_patterns = Counter(
        (row[0], row[3], row[4]) for row in data_mismatches
    )
    observations = []
    if total_missing:
        observations.append([
            "High", "Missing load coverage", total_missing,
            "Trace source records through staging, rejection logs, and load filters.",
        ])
    if total_unexpected:
        observations.append([
            "High", "Unexpected destination records", total_unexpected,
            "Review stale data, scope filters, and repeated loads.",
        ])
    if total_duplicate_rows:
        observations.append([
            "High", "Duplicate key values", total_duplicate_rows,
            "Confirm the Key Sequence definition and investigate duplicate business keys.",
        ])
    for (source_sheet, source_column, destination_column), count in mismatch_patterns.most_common(20):
        observations.append([
            "High" if count >= 100 else "Medium",
            f"{source_sheet}: {source_column} -> {destination_column}",
            count,
            "Review field mapping, transformation, datatype, trimming, and field-length rules.",
        ])
    if not observations:
        observations.append([
            "Low", "No exceptions detected", 0,
            "Retain the report as reconciliation evidence.",
        ])
    append_table(
        report, "Observations and Root Cause",
        ["Severity", "Pattern", "Exception Count", "Root Cause / Next Check"],
        observations,
    )

    recommendations = []
    if status == "PASS":
        recommendations.append([1, "No corrective action required", "Retain the report for sign-off."])
    else:
        recommendations.extend([
            [1, "Correct mapping and transformation issues", "Prioritize high-volume mismatch patterns."],
            [2, "Resolve missing and unexpected records", "Review scope, rejection, purge, and delta-load logic."],
            [3, "Resolve duplicate keys", "Correct data duplicates or revise Key Sequence metadata if incomplete."],
            [4, "Rerun reconciliation", "Require zero unexplained exceptions before final sign-off."],
        ])
    append_table(
        report, "Recommendations",
        ["Priority", "Corrective Action", "Implementation Guidance"],
        recommendations,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report.save(output_path)
    normalize_xlsx_font_order(output_path)
    print(f"Created: {output_path}")


def input_workbooks(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    if path.is_dir():
        return sorted(
            file for file in path.glob("*.xlsx")
            if not file.name.startswith("~$") and not file.name.endswith(OUTPUT_SUFFIX)
        )
    raise FileNotFoundError(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate one workbook or every .xlsx workbook in a folder using Mapping-Sheet metadata."
    )
    parser.add_argument("input", help="Input .xlsx file or folder")
    parser.add_argument("--output", help="Output report path; valid only for a single input file")
    parser.add_argument("--output-dir", help="Directory for reports in folder mode")
    parser.add_argument("--ignore-case", action="store_true", help="Compare text without case sensitivity")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    files = input_workbooks(input_path)
    if not files:
        print("No .xlsx files found", file=sys.stderr)
        return 2
    if args.output and len(files) != 1:
        parser.error("--output can only be used with one input workbook")

    failures = []
    for workbook_path in files:
        if args.output:
            output_path = Path(args.output).expanduser().resolve()
        else:
            output_directory = (
                Path(args.output_dir).expanduser().resolve()
                if args.output_dir else workbook_path.parent
            )
            output_path = output_directory / f"{workbook_path.stem}{OUTPUT_SUFFIX}"
        try:
            validate_workbook(workbook_path, output_path, args.ignore_case)
        except Exception as exc:
            failures.append((workbook_path, exc))
            print(f"FAILED: {workbook_path}: {exc}", file=sys.stderr)

    if failures:
        print(f"Completed with {len(failures)} failure(s).", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
